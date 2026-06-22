#!/usr/bin/env python3
"""Prepare search terms from the Huxiaoli script-name workbook."""

from __future__ import annotations

import argparse
import csv
import os
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT = SCRIPT_DIR / "狐小狸剧本杀.xlsx"
DEFAULT_OUTPUT = SCRIPT_DIR / "huxiaoli_search_terms.csv"
DEFAULT_VIEW_OUTPUT = SCRIPT_DIR / "huxiaoli_search_terms_for_viewing.csv"
DEFAULT_RATINGS_OUTPUT = SCRIPT_DIR / "huxiaoli_ipad_ratings.csv"
DEFAULT_REVIEWS_OUTPUT = SCRIPT_DIR / "huxiaoli_ipad_reviews.csv"
DEFAULT_DEBUG_DIR = SCRIPT_DIR / "debug_huxiaoli_ipad"
DEFAULT_SKIP_RATINGS_FROM = [
    SCRIPT_DIR / "miquan_ipad_ratings.csv",
    SCRIPT_DIR / "miquan_ratings.csv",
]

BRACKET_PATTERNS = [
    # These are catalog annotations, not part of the title to search in Miquan.
    re.compile(r"（[^（）]*）"),
    re.compile(r"\([^()]*\)"),
    re.compile(r"【[^【】]*】"),
    re.compile(r"\[[^\[\]]*\]"),
]

TRAILING_GENRE_OR_META = (
    # Tokens commonly appended after the real title, e.g.
    # "余香（5人）硬核还原推理情感" should search as "余香".
    "硬核",
    "还原",
    "推理",
    "情感",
    "沉浸",
    "欢乐",
    "本格",
    "阵营",
    "古风",
    "仙侠",
    "现代",
    "日式",
    "欧式",
    "欧美",
    "欧洲",
    "民国",
    "机制",
    "反转",
    "烧脑",
    "新手",
    "进阶",
    "开放",
    "封闭",
    "城限",
    "盒装",
    "独家",
    "恐怖",
    "惊悚",
    "悬疑",
    "立意",
    "治愈",
    "温情",
    "谍战",
    "玄幻",
    "变格",
    "架空",
    "科幻",
    "末世",
    "武侠",
    "校园",
    "都市",
    "魔幻",
    "跑团",
    "演绎",
    "微恐",
    "本",
)


def collapse_spaces(text: str) -> str:
    """Normalize all whitespace runs to a single regular space."""
    return re.sub(r"\s+", " ", text).strip()


def strip_bracketed_text(text: str) -> str:
    """Remove bracketed workbook annotations such as player counts or file notes."""
    previous = None
    while previous != text:
        previous = text
        for pattern in BRACKET_PATTERNS:
            text = pattern.sub("", text)
    return text


def strip_leading_catalog_marker(text: str) -> str:
    """Remove workbook sorting prefixes like 'A ' or numeric catalog IDs."""
    text = re.sub(r"^[A-ZＡ-Ｚ]\s+", "", text)
    text = re.sub(r"^\d+[\s.、_\-]+", "", text)
    return text


def strip_wrapping_quotes(text: str) -> str:
    """Remove quote/book-title wrappers only when they wrap the whole title."""
    pairs = (("《", "》"), ("「", "」"), ("『", "』"), ("“", "”"), ('"', '"'), ("'", "'"))
    changed = True
    while changed:
        changed = False
        text = text.strip()
        for left, right in pairs:
            if text.startswith(left) and text.endswith(right) and len(text) > len(left) + len(right):
                text = text[len(left) : -len(right)]
                changed = True
    return text


def strip_trailing_metadata(text: str) -> str:
    """Peel off trailing people-count and genre words until the title stabilizes."""
    previous = None
    while previous != text:
        previous = text
        # Keep meaningful title numbers like "来电2"; only strip counts ending in 人/男女.
        text = re.sub(r"\d+\s*[-|至]?\s*\d*\s*人$", "", text)
        text = re.sub(r"\d+\s*[男女]\s*\d*\s*[男女]$", "", text)
        for token in TRAILING_GENRE_OR_META:
            if text.endswith(token) and len(text) > len(token):
                text = text[: -len(token)]
                break
        # Avoid artifacts like "七个密室密室" after removing "烧脑推理本格".
        text = re.sub(r"(.+)([\u4e00-\u9fff]{2})\2$", r"\1\2", text)
        text = collapse_spaces(text)
    return text


def clean_script_name(raw_name: object) -> str:
    """Convert one workbook first-column value into the query used by the scraper."""
    text = "" if raw_name is None else str(raw_name)
    # NFKC converts full-width Latin letters/numbers and punctuation to standard forms.
    text = unicodedata.normalize("NFKC", text)
    text = collapse_spaces(text)
    text = strip_leading_catalog_marker(text)
    text = strip_bracketed_text(text)
    text = strip_wrapping_quotes(text)
    text = strip_trailing_metadata(text)
    text = re.sub(r"\s+", "", text)
    return text.strip(" -_、，,.;；:：")


def dedupe_key(text: str) -> str:
    """Build a comparison key so equivalent cleaned titles are only scraped once."""
    return "".join(ch for ch in text.casefold() if ch.isalnum() or "\u4e00" <= ch <= "\u9fff")


def workbook_rows(path: Path, sheet_names: list[str] | None = None) -> list[dict[str, str]]:
    """Read selected workbook sheets and return deduped scraper rows with audit data."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    selected = sheet_names or workbook.sheetnames
    rows: list[dict[str, str]] = []
    seen: set[str] = set()

    for sheet_name in selected:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}")
        sheet = workbook[sheet_name]
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            # Column A is the noisy title. Other columns are kept only for traceability.
            original = row[0] if row else None
            query = clean_script_name(original)
            if not query:
                continue
            key = dedupe_key(query)
            if key in seen:
                continue
            seen.add(key)
            rows.append(
                {
                    "query": query,
                    "original_name": "" if original is None else str(original).strip(),
                    "sheet": sheet_name,
                    "row": str(row_number),
                    "source_link": "" if len(row) < 2 or row[1] is None else str(row[1]).strip(),
                    "people_count": "" if len(row) < 3 or row[2] is None else str(row[2]).strip(),
                }
            )
    return rows


def write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    """Write cleaned terms in the CSV shape expected by the Appium scraper."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = ["query", "original_name", "sheet", "row", "source_link", "people_count"]
    # The scraper reads with plain utf-8 and expects the first header to be exactly "query".
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def write_view_csv(path: Path, rows: list[dict[str, str]]) -> None:
    """Write a BOM-marked CSV that opens with readable Chinese in spreadsheet apps."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = ["query", "original_name", "sheet", "row", "source_link", "people_count"]
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def scrape_processed_terms(args: argparse.Namespace, prepared_csv: Path) -> None:
    """Run the existing iPad scraper against the cleaned Huxiaoli query CSV."""
    # Import lazily so this file can still preprocess XLSX files without Appium installed.
    import miquan_appium_scraper as base
    import miquan_ipad_appium_scraper as ipad

    ipad.configure_base_for_ipad(args)
    if args.no_skip_existing_ratings:
        base.read_queries = ipad.ORIGINAL_READ_QUERIES
    else:
        ipad.install_completed_query_filter(args.skip_ratings_from)

    base.scrape(
        prepared_csv,
        args.reviews_output,
        args.ratings_output,
        args.mode,
        args.max_scrolls,
        args.reviews_per_script,
        None,
        args.debug_dir,
        args.rerun_existing,
    )


def main() -> None:
    """Parse CLI arguments, clean the workbook, and optionally scrape on iPad."""
    parser = argparse.ArgumentParser(description="Clean Huxiaoli workbook names into scraper search terms.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--view-output",
        type=Path,
        default=DEFAULT_VIEW_OUTPUT,
        help="Excel-friendly UTF-8-BOM copy for viewing Chinese correctly. The scraper does not use this file.",
    )
    parser.add_argument(
        "--no-view-output",
        action="store_true",
        help="Do not write the Excel-friendly viewing CSV.",
    )
    parser.add_argument(
        "--sheet",
        action="append",
        dest="sheets",
        help="Sheet to process. Repeat for multiple sheets. Defaults to every sheet with duplicate queries removed.",
    )
    parser.add_argument("--limit", type=int, default=None, help="Write only the first N cleaned terms.")
    parser.add_argument("--scrape", action="store_true", help="Run the iPad scraper after writing the cleaned CSV.")
    parser.add_argument(
        "--mode",
        choices=("ratings", "reviews", "both"),
        default="ratings",
        help="Scrape mode used only with --scrape.",
    )
    parser.add_argument("--ratings-output", type=Path, default=DEFAULT_RATINGS_OUTPUT)
    parser.add_argument("--reviews-output", type=Path, default=DEFAULT_REVIEWS_OUTPUT)
    parser.add_argument("--debug-dir", type=Path, default=DEFAULT_DEBUG_DIR)
    parser.add_argument("--max-scrolls", type=int, default=12)
    parser.add_argument("--reviews-per-script", type=int, default=10)
    parser.add_argument(
        "--rerun-existing",
        action="store_true",
        help="With --scrape, ignore existing rows in the Huxiaoli ratings output.",
    )
    parser.add_argument(
        "--skip-ratings-from",
        type=Path,
        action="append",
        default=None,
        help="With --scrape, skip queries already rated in this CSV. Repeat for multiple files.",
    )
    parser.add_argument(
        "--no-skip-existing-ratings",
        action="store_true",
        help="With --scrape, do not skip titles already rated in previous Miquan CSVs.",
    )
    parser.add_argument("--appium-url", default=os.environ.get("APPIUM_URL", "http://127.0.0.1:4723"))
    parser.add_argument("--device-name", default=os.environ.get("MIQUAN_IPAD_DEVICE_NAME", "iPad (2)"))
    parser.add_argument("--udid", default=os.environ.get("MIQUAN_IPAD_UDID", "00008112-0019593922DBA01E"))
    parser.add_argument("--platform-version", default=os.environ.get("MIQUAN_IPAD_PLATFORM_VERSION", "26.3.1"))
    parser.add_argument("--wda-port", type=int, default=int(os.environ.get("MIQUAN_IPAD_WDA_PORT", "8102")))
    parser.add_argument(
        "--wda-bundle-id",
        default=os.environ.get("MIQUAN_IPAD_WDA_BUNDLE_ID", "com.guozhan.WebDriverAgentRunner"),
    )
    parser.add_argument(
        "--use-new-wda",
        action="store_true",
        help="With --scrape, force Appium to rebuild/reinstall WebDriverAgent.",
    )
    args = parser.parse_args()
    if args.skip_ratings_from is None:
        args.skip_ratings_from = DEFAULT_SKIP_RATINGS_FROM

    rows = workbook_rows(args.input, args.sheets)
    if args.limit is not None:
        rows = rows[: args.limit]
    write_csv(args.output, rows)
    print(f"Wrote {len(rows)} search terms to {args.output}")
    if not args.no_view_output:
        write_view_csv(args.view_output, rows)
        print(f"Wrote Excel-friendly viewing copy to {args.view_output}")
    if args.scrape:
        scrape_processed_terms(args, args.output)


if __name__ == "__main__":
    main()
